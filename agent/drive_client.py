"""Google Drive access for the Portfolio Agent.

Two OAuth paths, sharing the same SCOPES:
- CLI (agent/portfolio_agent.py::main): a local installed-app flow
  (_get_credentials) that opens a browser on this machine and writes one
  shared token.json - unchanged from before the web app existed.
- Web app (agent/web.py): a proper web OAuth flow (build_auth_flow) where
  each visitor signs in with their own Google account; their credentials
  are stored in their own Flask session, never in a shared file, so
  concurrent users each see only their own Drive. get_service() picks
  whichever path applies based on whether it's called inside a Flask
  request context.

Almost everything here only reads from Drive (drive.readonly). The one
exception is upload_docx_as_google_doc (see below), which needs drive.file -
access limited to files this app itself creates, not arbitrary read/write to
the user's whole Drive - to create the Google Docs export. A session/token
issued before drive.file was added to SCOPES won't have it; see
upload_docx_as_google_doc's docstring for how that failure surfaces.
"""

import io
import json
import os
from pathlib import Path

import zipfile

import docx
from docx.opc.exceptions import PackageNotFoundError
from flask import g, has_request_context, session
from google.auth.transport.requests import Request
from google.oauth2.credentials import Credentials
from google_auth_oauthlib.flow import Flow, InstalledAppFlow
from googleapiclient.discovery import build
from googleapiclient.http import MediaIoBaseDownload, MediaIoBaseUpload
from pypdf import PdfReader
from pypdf.errors import PdfReadError

SCOPES = [
    "https://www.googleapis.com/auth/drive.readonly",
    "https://www.googleapis.com/auth/drive.file",
]

REPO_ROOT = Path(__file__).resolve().parent.parent
TOKEN_PATH = REPO_ROOT / "token.json"

GOOGLE_DOC_MIME_TYPE = "application/vnd.google-apps.document"
DOCX_MIME_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

# Google Workspace files can't be downloaded directly - they must be exported
# to a plain format first.
GOOGLE_EXPORT_MIME_TYPES = {
    GOOGLE_DOC_MIME_TYPE: "text/plain",
    "application/vnd.google-apps.spreadsheet": "text/csv",
    "application/vnd.google-apps.presentation": "text/plain",
}

# File types we can safely decode as text after downloading.
TEXT_MIME_PREFIXES = ("text/",)
TEXT_MIME_EXACT = {"application/json"}

FOLDER_MIME_TYPE = "application/vnd.google-apps.folder"

# PDFs have no native "read this document" input on Groq's chat completions
# API (unlike Gemini's multimodal document blocks), so text is extracted
# locally with pypdf and treated like any other text file from then on.
DOCUMENT_MIME_TYPES = {"application/pdf"}

# .docx (the modern, XML-based Word format) is handled the same way, via
# python-docx. The legacy binary .doc format (application/msword) isn't
# supported - python-docx can't read it - so it still falls through to
# "unreadable" below, same as before.
WORD_MIME_TYPES = {DOCX_MIME_TYPE}

# Images are intentionally not extracted - this agent's job is project text
# content (briefs, research notes, decisions), not visual analysis. Skipping
# them avoids the extra download/attachment work and the token cost of
# sending images to the model.
IMAGE_MIME_PREFIX = "image/"

# Must exactly match an "Authorized redirect URI" on the Web application
# OAuth client in Google Cloud Console (trailing slash included).
LOCAL_REDIRECT_PORT = 8080
LOCAL_REDIRECT_URI = f"http://localhost:{LOCAL_REDIRECT_PORT}/"


def _load_client_config():
    client_id = os.environ.get("CLIENT_ID")
    client_secret = os.environ.get("CLIENT_SECRET")
    if not client_id or not client_secret:
        raise RuntimeError(
            "CLIENT_ID and CLIENT_SECRET must be set in the environment "
            "(see .env) to authenticate with Google Drive."
        )
    return {
        "web": {
            "client_id": client_id,
            "client_secret": client_secret,
            "auth_uri": "https://accounts.google.com/o/oauth2/auth",
            "token_uri": "https://oauth2.googleapis.com/token",
            "redirect_uris": [LOCAL_REDIRECT_URI],
        }
    }


def _get_credentials():
    creds = None
    if TOKEN_PATH.exists():
        creds = Credentials.from_authorized_user_info(
            json.loads(TOKEN_PATH.read_text()), SCOPES
        )

    if creds and creds.expired and creds.refresh_token:
        creds.refresh(Request())

    if not creds or not creds.valid:
        flow = InstalledAppFlow.from_client_config(_load_client_config(), SCOPES)
        creds = flow.run_local_server(port=LOCAL_REDIRECT_PORT)

    TOKEN_PATH.write_text(creds.to_json())
    return creds


def build_auth_flow(redirect_uri: str, code_verifier: str | None = None) -> Flow:
    """Build a web OAuth Flow for the given callback URL (see agent/web.py's
    /auth/login and /auth/callback routes). `redirect_uri` must exactly
    match an Authorized redirect URI registered on this OAuth client in
    Google Cloud Console.

    PKCE note: Flow generates its own `code_verifier` the first time
    `authorization_url()` is called and stores it only on that Flow
    instance - it's never sent to Google as part of the redirect (only its
    hashed `code_challenge` is). Since /auth/login and /auth/callback are
    separate requests, each building a fresh Flow, the verifier has to be
    threaded through explicitly (via the session) or `fetch_token()` fails
    with "Missing code verifier". Callers must pass back whatever
    `flow.code_verifier` was after the /auth/login call.
    """
    flow = Flow.from_client_config(_load_client_config(), SCOPES, code_verifier=code_verifier)
    flow.redirect_uri = redirect_uri
    return flow


def credentials_to_dict(creds: Credentials) -> dict:
    """Serialize credentials for storage in a Flask session (JSON-safe)."""
    return {
        "token": creds.token,
        "refresh_token": creds.refresh_token,
        "token_uri": creds.token_uri,
        "client_id": creds.client_id,
        "client_secret": creds.client_secret,
        "scopes": creds.scopes,
    }


def credentials_from_dict(data: dict) -> Credentials:
    return Credentials(
        token=data.get("token"),
        refresh_token=data.get("refresh_token"),
        token_uri=data.get("token_uri"),
        client_id=data.get("client_id"),
        client_secret=data.get("client_secret"),
        scopes=data.get("scopes"),
    )


def _credentials_from_session() -> Credentials:
    """Load this web visitor's own Drive credentials from their Flask
    session, refreshing (and writing back) an expired access token via its
    refresh token. Raises RuntimeError if they haven't signed in."""
    data = session.get("credentials")
    if not data:
        raise RuntimeError("Not signed in - no Drive credentials in this session.")
    creds = credentials_from_dict(data)
    if creds.expired and creds.refresh_token:
        creds.refresh(Request())
        session["credentials"] = credentials_to_dict(creds)
    return creds


_service = None


def get_service():
    """Return a Drive API service. Inside a Flask request (the web app),
    this is scoped to the current request/visitor via flask.g, built from
    their own session credentials - never shared across users. Outside a
    request context (the CLI), falls back to the single local token.json
    flow, unchanged."""
    if has_request_context():
        if not hasattr(g, "_drive_service"):
            g._drive_service = build("drive", "v3", credentials=_credentials_from_session())
        return g._drive_service

    global _service
    if _service is None:
        _service = build("drive", "v3", credentials=_get_credentials())
    return _service


# --- Write access (drive.file scope) -------------------------------------
# Everything above this point only ever reads from Drive (drive.readonly).
# This is the one write path: exporting a finalized case study as a new
# Google Doc in the signed-in user's own Drive (see agent/web.py's
# /api/export_doc). It only needs access to files this app itself creates,
# which is exactly what the drive.file scope grants (see SCOPES above).


def upload_docx_as_google_doc(name: str, docx_bytes: bytes) -> dict:
    """Upload `docx_bytes` (see portfolio_agent.build_case_study_docx) as a
    new file in the signed-in user's own Drive, requesting that Drive convert
    it into a native, editable Google Doc on upload - passing
    mimeType=GOOGLE_DOC_MIME_TYPE in the create body while uploading .docx
    media is what triggers that conversion, rather than just storing the raw
    .docx file as-is. Word's heading styles carry over onto Google Docs'
    heading styles, so build_case_study_docx's add_heading calls produce a
    properly outlined Doc, not flat paragraphs.

    Requires the drive.file scope (see SCOPES). A token issued before that
    scope was added will fail this call with an HttpError whose status is 403
    and whose reason mentions insufficient scopes - callers should catch that
    specifically and tell the user to sign out and back in (see
    agent/web.py's /api/export_doc, which does this).

    Returns the created file's Drive API resource, e.g.
    {"id": "...", "webViewLink": "https://docs.google.com/document/d/.../edit"}.
    """
    service = get_service()
    media = MediaIoBaseUpload(io.BytesIO(docx_bytes), mimetype=DOCX_MIME_TYPE, resumable=False)
    file_metadata = {"name": name, "mimeType": GOOGLE_DOC_MIME_TYPE}
    return (
        service.files()
        .create(body=file_metadata, media_body=media, fields="id, webViewLink")
        .execute()
    )


def search_items(query: str, parent_id: str | None = None) -> list[dict]:
    """Search Drive for files/folders whose name contains `query`."""
    service = get_service()
    safe_query = query.replace("'", "\\'")
    q = f"name contains '{safe_query}' and trashed = false"
    if parent_id:
        q += f" and '{parent_id}' in parents"

    results = (
        service.files()
        .list(
            q=q,
            fields="files(id, name, mimeType, parents, modifiedTime)",
            pageSize=25,
            includeItemsFromAllDrives=True,
            supportsAllDrives=True,
        )
        .execute()
    )
    return results.get("files", [])


def list_folder(folder_id: str) -> list[dict]:
    """List the immediate contents of a Drive folder."""
    service = get_service()
    q = f"'{folder_id}' in parents and trashed = false"
    results = (
        service.files()
        .list(
            q=q,
            fields="files(id, name, mimeType, modifiedTime)",
            pageSize=100,
            includeItemsFromAllDrives=True,
            supportsAllDrives=True,
        )
        .execute()
    )
    return results.get("files", [])


def get_file_metadata(file_id: str) -> dict:
    service = get_service()
    return (
        service.files()
        .get(fileId=file_id, fields="id, name, mimeType, modifiedTime, size")
        .execute()
    )


def _download_bytes(file_id: str) -> bytes:
    service = get_service()
    request = service.files().get_media(fileId=file_id)
    buffer = io.BytesIO()
    downloader = MediaIoBaseDownload(buffer, request)
    done = False
    while not done:
        _, done = downloader.next_chunk()
    return buffer.getvalue()


def _extract_pdf_text(raw: bytes) -> str:
    reader = PdfReader(io.BytesIO(raw))
    return "\n\n".join(page.extract_text() or "" for page in reader.pages).strip()


def _extract_docx_text(raw: bytes) -> str:
    document = docx.Document(io.BytesIO(raw))
    parts = [p.text for p in document.paragraphs if p.text]
    for table in document.tables:
        for row in table.rows:
            row_text = " | ".join(cell.text for cell in row.cells if cell.text)
            if row_text:
                parts.append(row_text)
    return "\n\n".join(parts).strip()


def read_file(file_id: str) -> dict:
    """Read a Drive file. Returns one of:
    - {"kind": "text", "text": str}
      (includes PDFs and .docx - text extracted locally with pypdf/python-docx,
      since the LLM has no native "read this document" input)
    - {"kind": "unreadable", "message": str}
      (includes images - intentionally not extracted, see IMAGE_MIME_PREFIX above)
    """
    service = get_service()
    metadata = get_file_metadata(file_id)
    mime_type = metadata.get("mimeType", "")
    name = metadata.get("name", file_id)

    if mime_type == FOLDER_MIME_TYPE:
        return {
            "kind": "unreadable",
            "message": f"'{name}' is a folder, not a file. Use list_folder to see its contents.",
        }

    if mime_type.startswith(IMAGE_MIME_PREFIX):
        return {
            "kind": "unreadable",
            "message": f"'{name}' is an image - not extracted (this agent only reads text/PDF project content).",
        }

    export_mime_type = GOOGLE_EXPORT_MIME_TYPES.get(mime_type)
    try:
        if export_mime_type:
            data = service.files().export(fileId=file_id, mimeType=export_mime_type).execute()
            text = data.decode("utf-8") if isinstance(data, bytes) else data
            return {"kind": "text", "text": text}

        if mime_type in DOCUMENT_MIME_TYPES:
            raw = _download_bytes(file_id)
            try:
                text = _extract_pdf_text(raw)
            except PdfReadError as exc:
                return {"kind": "unreadable", "message": f"Could not parse '{name}' as a PDF: {exc}"}
            if not text:
                return {"kind": "unreadable", "message": f"'{name}' is a PDF with no extractable text (scanned/image-only page?)."}
            return {"kind": "text", "text": text}

        if mime_type in WORD_MIME_TYPES:
            raw = _download_bytes(file_id)
            try:
                text = _extract_docx_text(raw)
            except (PackageNotFoundError, zipfile.BadZipFile) as exc:
                return {"kind": "unreadable", "message": f"Could not parse '{name}' as a .docx file: {exc}"}
            if not text:
                return {"kind": "unreadable", "message": f"'{name}' is a .docx file with no extractable text."}
            return {"kind": "text", "text": text}

        if mime_type.startswith(TEXT_MIME_PREFIXES) or mime_type in TEXT_MIME_EXACT:
            raw = _download_bytes(file_id)
            return {"kind": "text", "text": raw.decode("utf-8", errors="replace")}
    except Exception as exc:  # noqa: BLE001 - surfaced to the agent as a tool result
        return {"kind": "unreadable", "message": f"Error reading '{name}' ({mime_type}): {exc}"}

    return {
        "kind": "unreadable",
        "message": (
            f"'{name}' has type '{mime_type}', which can't be extracted as text or read "
            "as a document. It was found but not read."
        ),
    }
